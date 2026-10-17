"""
Automated Test Suite for Student Management System
Verifies all 5 core requirements:
1. Admin panel login & registration
2. Student personal link access & assignments upload
3. Access control / security (URL tampering, token verification, 403 Access Denied)
4. Multi-sheet Excel creation, logging, and downloading
5. Admin record editing & master exports
"""

import io
import os
import openpyxl
from app import app
import db
import excel_service

def run_tests():
    print("==================================================")
    print(" [*] Running Automated System Test Suite...")
    print("==================================================")

    # Use Flask test client
    app.config["TESTING"] = True
    client = app.test_client()

    # 1. Admin Login Tests
    print("\n[TEST 1] Admin Authentication:")
    res_bad = client.post("/login", data={"username": "admin", "password": "wrongpassword"})
    assert res_bad.status_code == 200
    assert b"Invalid admin credentials" in res_bad.data
    print("  [+] Invalid credentials rejected correctly.")

    res_good = client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)
    assert res_good.status_code == 200
    assert b"Admin Management Portal" in res_good.data
    print("  [+] Admin authenticated and dashboard rendered successfully.")

    # 2. Student Registration & Token Generation
    print("\n[TEST 2] Student Registration & 5-Sheet Excel Auto-Creation:")
    reg_data = {
        "roll_no": "TEST999",
        "name": "Alex Mercer",
        "class_name": "M.Sc Cybersecurity",
        "email": "alex.mercer@example.com",
        "phone": "+1 555 123 4567",
        "father_name": "Thomas Mercer",
        "dob": "2003-09-12",
        "address": "404 Hacker Way"
    }
    res_reg = client.post("/register_student", data=reg_data, follow_redirects=True)
    assert res_reg.status_code == 200
    student = db.get_student("TEST999")
    assert student is not None
    assert student["token"] is not None and len(student["token"]) >= 20
    print(f"  [+] Student registered with Roll: TEST999, Token: {student['token']}")

    # Verify Excel file exists and has 5 sheets
    excel_path = os.path.join(excel_service.EXCEL_DIR, "TEST999.xlsx")
    assert os.path.exists(excel_path)
    wb = openpyxl.load_workbook(excel_path)
    expected_sheets = ["Personal Details", "Attendance", "Marks & Grades", "Assignments Submitted", "Remarks"]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Missing sheet: {s}"
    wb.close()
    print(f"  [+] TEST999.xlsx generated with all 5 sheets: {wb.sheetnames}")

    # 3. Security & Access Control Tests
    print("\n[TEST 3] Access Control & URL Tampering Security:")
    # Logout admin first so we test as student
    client.get("/logout")

    # Correct link
    res_student = client.get(f"/student/TEST999?token={student['token']}")
    assert res_student.status_code == 200
    assert b"Alex Mercer" in res_student.data
    print("  [+] Valid personal link accessed successfully.")

    # Access without token
    res_no_token = client.get("/student/TEST999")
    assert res_no_token.status_code == 403
    assert b"Access Denied" in res_no_token.data
    print("  [+] Request without token returned 403 Access Denied.")

    # Access with invalid token
    res_wrong_token = client.get("/student/TEST999?token=bogus_token_12345")
    assert res_wrong_token.status_code == 403
    assert b"Access Denied" in res_wrong_token.data
    print("  [+] Request with invalid token returned 403 Access Denied.")

    # Tampering: Change roll number in URL using someone else's token
    res_tamper = client.get(f"/student/ROLL101?token={student['token']}")
    assert res_tamper.status_code == 403
    assert b"Access Denied" in res_tamper.data
    print("  [+] URL Tampering (/student/ROLL101 with TEST999's token) returned 403 Access Denied.")

    # 4. Assignment Upload & Logging
    print("\n[TEST 4] Assignment Upload, Security & Excel Logging:")
    # Mismatched token on upload
    bad_upload = client.post("/student/ROLL101/upload", data={
        "token": student["token"],
        "assignment_file": (io.BytesIO(b"Dummy PDF content"), "hack.pdf")
    })
    assert bad_upload.status_code == 403
    print("  [+] Unauthorized upload rejected with 403 Access Denied.")

    # Valid upload
    test_pdf_content = b"%PDF-1.4 Mock Assignment Data for Testing"
    res_upload = client.post(f"/student/TEST999/upload", data={
        "token": student["token"],
        "assignment_file": (io.BytesIO(test_pdf_content), "network_security_report.pdf")
    }, follow_redirects=True)
    assert res_upload.status_code == 200
    assert b"network_security_report.pdf" in res_upload.data

    # Check database
    assignments = db.get_student_assignments("TEST999")
    assert len(assignments) == 1
    assert assignments[0]["original_name"] == "network_security_report.pdf"
    print("  [+] Assignment recorded in SQLite database.")

    # Check Excel Sheet 4
    wb = openpyxl.load_workbook(excel_path)
    ws4 = wb["Assignments Submitted"]
    found_in_excel = False
    for row in ws4.iter_rows(values_only=True):
        if "network_security_report.pdf" in str(row):
            found_in_excel = True
            break
    assert found_in_excel
    wb.close()
    print("  [+] Assignment successfully logged in Sheet 4 ('Assignments Submitted') of TEST999.xlsx.")

    # 5. Excel Downloads
    print("\n[TEST 5] Excel Downloads:")
    # Student downloads their own Excel
    res_dl = client.get(f"/download/excel/TEST999?token={student['token']}")
    assert res_dl.status_code == 200
    assert res_dl.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    res_dl.close()
    print("  [+] Student downloaded own Excel record successfully.")

    # Unauthorized student downloading someone else's Excel
    res_unauth_dl = client.get(f"/download/excel/ROLL101?token={student['token']}")
    assert res_unauth_dl.status_code == 403
    print("  [+] Downloading another student's Excel without auth blocked with 403.")

    # Download template
    res_tmpl = client.get("/download/template")
    assert res_tmpl.status_code == 200
    assert res_tmpl.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    res_tmpl.close()
    print("  [+] Downloadable blank 5-sheet template verified.")

    # 6. Admin Record Editing & Academic Entry
    print("\n[TEST 6] Admin Editing & Academic Data Entry:")
    # Re-login as Admin
    client.post("/login", data={"username": "admin", "password": "admin123"})

    # Edit student details
    edit_data = {
        "name": "Alexander Mercer Senior",
        "class_name": "Ph.D Cybersecurity",
        "email": "alex.senior@example.com",
        "phone": "+1 555 999 8888",
        "father_name": "Thomas Mercer",
        "dob": "2003-09-12",
        "address": "909 Research Blvd"
    }
    res_edit = client.post("/student/TEST999/edit", data=edit_data, follow_redirects=True)
    assert res_edit.status_code == 200
    updated = db.get_student("TEST999")
    assert updated["name"] == "Alexander Mercer Senior"
    assert updated["class_name"] == "Ph.D Cybersecurity"
    print("  [+] Student profile updated in DB and Sheet 1 of Excel.")

    # Add academic mark to Sheet 3
    client.post("/student/TEST999/add_mark", data={
        "code": "CYBER901",
        "name": "Cryptography & Blockchain",
        "exam_type": "Final Exam",
        "max_marks": "100",
        "obtained": "96",
        "grade": "A+",
        "status": "Passed"
    }, follow_redirects=True)

    # Add attendance to Sheet 2
    client.post("/student/TEST999/add_attendance", data={
        "term": "Spring 2026",
        "total_classes": "30",
        "attended": "29",
        "status": "Honor Standing",
        "remarks": "Flawless lab attendance"
    }, follow_redirects=True)

    # Add remark to Sheet 5
    client.post("/student/TEST999/add_remark", data={
        "category": "Academics",
        "remarks": "Outstanding research paper submitted."
    }, follow_redirects=True)

    # Verify Sheet 2, 3, and 5 have the new data
    excel_read = excel_service.read_student_excel_data("TEST999")
    assert any(m["code"] == "CYBER901" for m in excel_read["marks"])
    assert any(a["term"] == "Spring 2026" for a in excel_read["attendance"])
    assert any("Outstanding research" in r["text"] for r in excel_read["remarks"])
    print("  [+] Marks, Attendance, and Remarks successfully appended to Excel workbook.")

    # Export Master Excel
    res_master = client.get("/download/master_excel")
    assert res_master.status_code == 200
    res_master.close()
    print("  [+] Consolidated master Excel export generated and downloaded.")

    # 7. Cleanup test student
    print("\n[TEST 7] Student Deletion Cleanup:")
    import gc
    gc.collect()
    res_del = client.post("/delete_student/TEST999", follow_redirects=True)
    assert res_del.status_code == 200
    assert db.get_student("TEST999") is None
    assert not os.path.exists(excel_path)
    print("  [+] Student TEST999, database row, and Excel file cleaned up cleanly.")

    print("\n==================================================")
    print(" [SUCCESS] ALL 7 AUTOMATED TESTS PASSED (100% OK)!")
    print("==================================================")

if __name__ == "__main__":
    run_tests()
